"""Builds the SOC policy templates that ship in packs/attest/templates.

    python docs/dev/build-soc-templates.py

Writes one Word file per document, index.json (what the Templates tab lists)
and fields.json (what the product can fill in). Run it again after changing a
document below; it replaces everything it made.

The shapes are the ones apps/api/src/templates/fill.ts expects:
  - a blank is a yellow <placeholder> run inside a content control whose tag
    is the blank's name. {name} is a shared answer (asked once, used in every
    document), {$owner} and friends come from the Policies screen, and
    [[NN.bN]] is a question of this document's own;
  - a table filled from a register has the caption "source:<register>" and one
    example row, copied once per record;
  - everything before the first page break is instructions, dropped when the
    product fills the document in.

All wording is original to Offset Security. Nothing here is the AICPA's text.
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parents[2] / "packs" / "attest" / "templates"
NAVY = RGBColor(0x0F, 0x1C, 0x36)

# ── answers asked once ───────────────────────────────────────────────────────
GROUPS = [
    {"id": "company", "title": "About your company",
     "note": "Asked once and put into every document. Leave anything you do not know yet; it stays yellow."},
    {"id": "choices", "title": "Your choices",
     "note": "The usual answers are already filled in. Change any that are not what you do: the policies must match what really happens, because the auditor tests what they say."},
]
# name: (group, placeholder, label, hint, default, long)
SHARED = {
    "company": ("company", "Company name", "Company name", "As it appears on contracts.", "", False),
    "system_name": ("company", "name of the service", "The service your SOC report covers",
                    "As customers know it, for example Acme Payroll Cloud.", "", False),
    "security_owner": ("company", "security owner role", "Who owns security (job title)",
                       "For example Head of Security or CTO. Owns these documents unless Policies says otherwise.", "", False),
    "approver": ("company", "approver role", "Who approves the policies (job title)",
                 "Usually the CEO, or the board.", "", False),
    "security_contact": ("company", "security contact", "Where people report security problems",
                         "An email address or channel, for example security@acme.com.", "", False),
    "hosting_provider": ("company", "hosting provider", "Where the system is hosted",
                         "For example AWS eu-west-1, or Microsoft Azure UK South.", "", False),
    "doc_classification": ("choices", "classification", "Classification printed on each document", "", "Internal", False),
    "review_cycle": ("choices", "review cycle", "How often policies are reviewed", "", "at least once a year", False),
    "access_review": ("choices", "access review frequency", "How often user access is reviewed", "", "every quarter", False),
    "deprovision_time": ("choices", "time to remove access", "How soon a leaver's access is removed", "", "one business day", False),
    "patch_critical": ("choices", "critical patch time", "Time to fix a critical vulnerability", "", "14 days", False),
    "patch_high": ("choices", "high patch time", "Time to fix a high vulnerability", "", "30 days", False),
    "backup_frequency": ("choices", "backup frequency", "How often production data is backed up", "", "daily", False),
    "rto": ("choices", "recovery time objective", "Recovery time objective", "How long the service may be down after a disaster.", "8 hours", False),
    "rpo": ("choices", "recovery point objective", "Recovery point objective", "How much data, in time, you can afford to lose.", "24 hours", False),
    "training_cycle": ("choices", "training frequency", "How often everyone repeats security training", "", "every year", False),
    "vendor_review": ("choices", "vendor review frequency", "How often critical vendors are reviewed", "", "every year", False),
    "log_retention": ("choices", "log retention", "How long security logs are kept", "", "12 months", False),
    "customer_notice": ("choices", "customer notice time", "How soon affected customers are told of a breach", "", "72 hours", False),
    "data_after_contract": ("choices", "deletion period", "How soon customer data is deleted after a contract ends", "", "90 days", False),
}
SPECIAL = {"$owner": "owner", "$approver": "approver", "$version": "version",
           "$approval_date": "effective date", "$next_review": "next review", "$classification": "classification"}

TOKEN = re.compile(r"\{(\$?\w+)\}|\[\[([\w.]+)\]\]")


class Doc:
    def __init__(self, no, slug, title, ref, about, fields=None):
        self.no, self.title, self.ref, self.about = no, title, ref, about
        self.file = f"{no}_{slug}.docx"
        self.fields = fields or {}  # key -> dict(ph, label, hint, default, long)
        self.body = []  # (kind, payload, section)
        self.section = "General"

    def h(self, text):
        self.section = text
        self.body.append(("h", text, text))
        return self

    def p(self, text):
        self.body.append(("p", text, self.section))
        return self

    def ul(self, *items):
        for i in items:
            self.body.append(("li", i, self.section))
        return self

    def table(self, header, rows):
        self.body.append(("table", (header, rows), self.section))
        return self

    def source(self, name, header):
        self.body.append(("source", (name, header), self.section))
        return self


def f(ph, label, hint="", default="", long=False):
    return {"ph": ph, "label": label, "hint": hint, "default": default, "long": long}


# ── the documents ────────────────────────────────────────────────────────────
DOCS = []

d = Doc("01", "Information_Security_Policy", "Information Security Policy", "CC1, CC2, CC5",
        "The top-level policy: who is responsible, what every other policy covers, and how they are kept current.")
d.h("Purpose").p("This policy sets out how {company} protects the confidentiality, integrity and availability of {system_name} and the information customers trust us with. It is the policy every other security policy sits under.")
d.h("Scope").p("It applies to everyone who works for {company}, including employees, contractors and temporary staff, and to every system, application and record inside the boundary of {system_name}.")
d.h("Roles and responsibilities").ul(
    "The {approver} approves this policy and the policies under it, and reviews the security programme at least once a year.",
    "The {security_owner} owns the security programme: keeping policies current, running the risk assessment, and reporting to the {approver}.",
    "Managers make sure their teams know and follow these policies.",
    "Everyone follows these policies and reports security problems to {security_contact} straight away.")
d.h("Policy").ul(
    "We assess security risks at least once a year and whenever something significant changes.",
    "We choose controls that bring each significant risk to a level the {approver} accepts.",
    "We keep the policies listed below, give each an owner, and review each {review_cycle}.",
    "We monitor whether controls work, record exceptions, and fix them on time.",
    "We tell customers about our security commitments and about changes that affect them.")
d.h("The policies under this one").source("policies", ["Ref", "Policy", "Criteria", "Owner", "Version", "Next review", "Where it is kept"])
d.h("Exceptions").p("Any exception to a security policy is requested in writing, approved by the {security_owner}, given an end date and recorded.")
d.h("Enforcement").p("Breaking this policy may lead to disciplinary action, up to and including dismissal, under [[01.b1]].")
d.fields = {"01.b1": f("disciplinary process", "Which process handles breaches of policy?", "For example 'the disciplinary procedure in the Employee Handbook'.", "the disciplinary procedure in the Employee Handbook")}
DOCS.append(d)

d = Doc("02", "Code_of_Conduct_and_Acceptable_Use", "Code of Conduct and Acceptable Use Policy", "CC1.1, CC1.5, CC2.2",
        "How people are expected to behave, and how they may use company systems and data. Everyone signs it.")
d.h("Purpose").p("{company} expects everyone to act with integrity. This policy says what that means day to day, and how company systems and information may be used.")
d.h("Conduct").ul(
    "Act honestly and fairly with customers, colleagues and suppliers.",
    "Declare any conflict of interest to your manager.",
    "Never offer or accept bribes, or gifts above [[02.b1]].",
    "Protect customer information as if it were your own.",
    "Speak up about wrongdoing through {security_contact} or [[02.b2]], without fear of retaliation.")
d.h("Acceptable use").ul(
    "Use company systems for company work. Limited personal use is allowed if it does not interfere with work or break this policy.",
    "Keep your passwords to yourself, use the password manager, and turn on multi-factor authentication wherever it is offered.",
    "Lock your screen when you step away, and keep company laptops encrypted and up to date.",
    "Do not install unapproved software or connect unapproved devices to company systems.",
    "Do not copy customer data to personal devices, personal email or unapproved cloud services.",
    "Report lost devices and suspected phishing to {security_contact} straight away.")
d.h("Acknowledgement").p("Everyone reads and signs this policy when they join and {training_cycle} after that. Breaking it may lead to disciplinary action.")
d.table(["Name", "Signature", "Date"], [["<name>", "", "<date>"]])
d.fields = {
    "02.b1": f("gift limit", "Gift limit", "The value above which a gift must be declined or declared.", "a nominal value"),
    "02.b2": f("anonymous channel", "Anonymous reporting channel", "Where someone can raise a concern without giving their name.", "the anonymous reporting line"),
}
DOCS.append(d)

d = Doc("03", "Access_Control_Policy", "Access Control Policy", "CC6.1, CC6.2, CC6.3",
        "Who gets access to what, how it is approved, reviewed and removed, and how accounts are protected.")
d.h("Purpose").p("Access to {system_name} and the information in it is limited to people who need it for their job, approved before it is given, and removed when it is no longer needed.")
d.h("Accounts and authentication").ul(
    "Everyone has their own account. Shared accounts are not allowed except where a system offers nothing else, and then their use is logged.",
    "Sign-in to company systems goes through [[03.b1]] with multi-factor authentication.",
    "Passwords are at least [[03.b2]] characters and are kept in the company password manager.",
    "Production systems are reached only through [[03.b3]].")
d.h("Giving access").ul(
    "Access is requested through a ticket and approved by the person's manager and the system owner before it is given.",
    "Access is given by role, with the least privilege the job needs.",
    "Administrator access is limited to named people and reviewed separately.")
d.h("Changing and removing access").ul(
    "When someone changes role, access they no longer need is removed.",
    "When someone leaves, all their access is removed within {deprovision_time} of their last day.",
    "HR tells IT about every leaver before their last day.")
d.h("Access reviews").p("System owners review who has access to each in-scope system {access_review}. Access that is no longer needed is removed, and the review is kept as evidence with the tickets it raised.")
d.h("Physical access").p("Our production systems run at {hosting_provider}, whose data centre controls are covered by its own SOC report. Office access is by badge, and visitors are signed in and escorted.")
d.fields = {
    "03.b1": f("identity provider", "Which identity provider do people sign in with?", "For example Okta, Google Workspace or Microsoft Entra ID."),
    "03.b2": f("minimum length", "Minimum password length", "", "12"),
    "03.b3": f("admin access route", "How do administrators reach production?", "For example 'the VPN', 'AWS SSO' or 'the bastion host'."),
}
DOCS.append(d)

d = Doc("04", "Change_Management_Policy", "Change Management Policy", "CC8.1",
        "How changes to code, infrastructure and configuration are requested, reviewed, tested, approved and released.")
d.h("Purpose").p("Every change to {system_name} in production is planned, reviewed, tested and approved, so that changes do not cause outages, security gaps or errors in customers' data.")
d.h("Normal changes").ul(
    "Every change is tracked in [[04.b1]] and linked to its code or configuration.",
    "Code is kept in [[04.b2]]. Changes reach the main branch only through a pull request reviewed by someone other than the author.",
    "Automated tests run on every pull request and must pass before merging.",
    "Changes are tested in an environment separate from production.",
    "Releases to production go through the deployment pipeline. Developers do not change production by hand.")
d.h("Infrastructure and configuration").p("Infrastructure is defined as code where possible and changed through the same review. Changes made in a console are recorded in a ticket.")
d.h("Emergency changes").p("A change needed to restore service or fix a serious vulnerability may be released before review. It is reviewed and approved within [[04.b3]] and recorded as an emergency change.")
d.h("Segregation of duties").p("Nobody approves their own change. Access to deploy to production is limited to the pipeline and named administrators.")
d.fields = {
    "04.b1": f("ticketing system", "Where are changes tracked?", "For example Jira, Linear or GitHub Issues."),
    "04.b2": f("code repository", "Where is code kept?", "For example GitHub or GitLab."),
    "04.b3": f("review window", "How soon is an emergency change reviewed?", "", "one business day"),
}
DOCS.append(d)

d = Doc("05", "Operations_and_Vulnerability_Management", "Operations and Vulnerability Management Policy", "CC6.6, CC6.7, CC6.8, CC7.1, CC7.2",
        "Network protection, encryption, malware protection, vulnerability scanning, patching and monitoring.")
d.h("Network and encryption").ul(
    "Production networks deny traffic by default. Firewall and security group rules are reviewed {review_cycle}.",
    "Data travelling over public networks is encrypted with TLS 1.2 or later.",
    "Customer data is encrypted at rest. Keys are kept in [[05.b1]] and access to them is restricted.",
    "Company laptops use full-disk encryption, enforced through device management.")
d.h("Malware protection").p("Endpoint protection runs on every company laptop and server, is kept up to date, and alerts the {security_owner} on detection.")
d.h("Vulnerability management").ul(
    "Systems and code dependencies are scanned for vulnerabilities at least [[05.b2]].",
    "Critical vulnerabilities are fixed within {patch_critical} and high ones within {patch_high}. Anything that cannot be fixed in time is recorded as an exception with a plan.",
    "An independent penetration test is carried out at least once a year, and its findings are tracked to closure.")
d.h("Logging and monitoring").ul(
    "Security-relevant logs from in-scope systems are collected in [[05.b3]] and kept for {log_retention}.",
    "Alerts for suspicious activity go to the on-call engineer, who records how each one was handled.",
    "Capacity and availability are monitored, with alerts before limits are reached.")
d.fields = {
    "05.b1": f("key management service", "Where are encryption keys kept?", "For example AWS KMS or Azure Key Vault."),
    "05.b2": f("scan frequency", "How often are vulnerability scans run?", "", "monthly"),
    "05.b3": f("log platform", "Where are logs collected?", "For example Datadog, Splunk or CloudWatch."),
}
DOCS.append(d)

d = Doc("06", "Incident_Response_Plan", "Incident Response Plan", "CC7.3, CC7.4, CC7.5, CC2.3",
        "How security events are triaged, incidents handled, customers told, and lessons learned.")
d.h("Purpose").p("This plan makes sure security incidents affecting {system_name} are found quickly, contained, and learned from, and that customers are told when they need to be.")
d.h("Reporting").p("Anyone who suspects a security incident reports it to {security_contact} immediately. Out of hours, page [[06.b1]].")
d.h("Severity").table(["Severity", "Meaning", "Respond within"], [
    ["Critical", "Customer data exposed, or the service down for all customers", "[[06.b2]]"],
    ["High", "A serious threat that could become critical, or a partial outage", "4 hours"],
    ["Medium", "A limited security event with no customer impact", "1 business day"],
    ["Low", "A policy breach or suspicious activity with no harm done", "5 business days"]])
d.h("Response").ul(
    "Triage: decide whether the event is an incident, set its severity and name an incident lead.",
    "Contain: stop the damage spreading, preserving evidence.",
    "Eradicate and recover: remove the cause, restore affected systems and confirm they are working.",
    "Communicate: keep the {approver} informed, and tell affected customers within {customer_notice} of confirming their data was affected.",
    "Learn: hold a review within two weeks, record what happened and why, and track the actions to closure.")
d.h("Testing").p("The plan is tested with a tabletop exercise at least once a year, and updated with what the exercise shows.")
d.h("Incident log").source("incidents", ["Incident", "Detected", "Severity", "Status", "Owner"])
d.fields = {
    "06.b1": f("on-call route", "How is someone reached out of hours?", "For example 'the on-call rota in PagerDuty'."),
    "06.b2": f("critical response time", "How soon must a critical incident get a response?", "", "1 hour"),
}
DOCS.append(d)

d = Doc("07", "Business_Continuity_and_Disaster_Recovery_Plan", "Business Continuity and Disaster Recovery Plan", "CC9.1, A1.1, A1.2, A1.3",
        "Recovery objectives, backups, redundancy, and how recovery is tested.")
d.h("Objectives").table(["Measure", "Target"], [
    ["Recovery time objective (RTO)", "{rto}"],
    ["Recovery point objective (RPO)", "{rpo}"],
    ["Availability commitment", "[[07.b1]]"]])
d.h("Resilience").ul(
    "{system_name} runs at {hosting_provider} across more than one availability zone.",
    "Production data is backed up {backup_frequency}, encrypted, and copied to a separate region. Backup failures alert the on-call engineer.",
    "Capacity is monitored and scaled ahead of demand.")
d.h("Recovery").ul(
    "The {security_owner} decides when to invoke this plan and leads the recovery.",
    "Recovery runbooks are kept in [[07.b2]], where the on-call team can reach them during an outage.",
    "Customers are kept informed through the status page.")
d.h("Testing").p("A restore from backup is tested at least every quarter, and the full disaster recovery plan at least once a year. Results, including the recovery time achieved, are recorded and gaps are fixed.")
d.h("Key people").ul("Business continuity also covers people: each critical role has a named deputy who can run it.")
d.fields = {
    "07.b1": f("availability target", "What availability do you promise customers?", "For example 99.9% monthly."),
    "07.b2": f("runbook location", "Where are the recovery runbooks?", "Somewhere that still works when production is down."),
}
DOCS.append(d)

d = Doc("08", "Risk_Assessment_Policy", "Risk Assessment Policy", "CC3.1, CC3.2, CC3.3, CC3.4",
        "How risks to the service commitments are found, scored, treated and reviewed, including fraud. The register is filled from Risks.")
d.h("Purpose").p("{company} assesses risks to meeting its service commitments at least once a year, so that controls are chosen for real risks and significant changes are not missed.")
d.h("Method").ul(
    "Risks are identified by the {security_owner} with the owners of each area, covering technology, people, vendors, processes and fraud.",
    "Each risk is scored for likelihood and impact from 1 to 5. The score is likelihood times impact.",
    "Risks scoring [[08.b1]] or more are treated: reduced with controls, transferred, avoided, or formally accepted by the {approver}.",
    "Each risk has an owner and a review date.")
d.h("Fraud").p("The assessment considers how someone could misuse access, alter records or override controls, alone or with others.")
d.h("Change").p("Significant changes, such as a new product, a new vendor that holds customer data, or a change of hosting, are assessed for risk before they happen.")
d.h("Risk register").source("risks", ["Ref", "Risk", "Score", "Treatment", "Owner", "Status"])
d.fields = {"08.b1": f("treatment threshold", "Which score must be treated?", "Risks at or above this score need a decision.", "12")}
DOCS.append(d)

d = Doc("09", "Vendor_Management_Policy", "Vendor Management Policy", "CC9.2, P6.4",
        "How vendors and subservice organisations are assessed, contracted and reviewed. The list is filled from Vendors.")
d.h("Purpose").p("Vendors that touch {system_name} or customer data are assessed before they are used and reviewed while they are, because their security is part of ours.")
d.h("Before using a vendor").ul(
    "The requester records the vendor, what it will do and what data it will hold.",
    "The {security_owner} rates the vendor's criticality and asks for assurance to match: a SOC 2 report or ISO 27001 certificate for critical vendors, a questionnaire for others.",
    "Contracts with vendors that hold customer data include confidentiality, security, breach notification within [[09.b1]], and deletion at the end of the contract.")
d.h("While using a vendor").ul(
    "Critical vendors are reviewed {vendor_review}: their latest SOC report is read, exceptions noted, and the controls they expect from us confirmed.",
    "Vendors on which our own controls rely are subservice organisations. Their SOC reports are reviewed and kept with our audit evidence.",
    "When a vendor is no longer used, its access is removed and deletion of our data is confirmed.")
d.h("Vendors").source("vendors", ["Service", "Vendor", "Assurance"])
d.fields = {"09.b1": f("vendor breach notice", "How soon must vendors tell you about a breach?", "", "72 hours")}
DOCS.append(d)

d = Doc("10", "Data_Classification_and_Retention_Policy", "Data Classification and Retention Policy", "C1.1, C1.2, CC6.5",
        "How information is classified and handled, how long it is kept, and how it is destroyed.")
d.h("Classification").table(["Class", "Examples", "Handling"], [
    ["Restricted", "Customer data, credentials, encryption keys", "Encrypted, access on a need-to-know basis, never on personal devices"],
    ["Confidential", "Contracts, source code, internal reports", "Company systems only, shared with care"],
    ["Internal", "Policies, procedures, internal messages", "Company systems only"],
    ["Public", "Website, marketing material", "No restrictions"]])
d.h("Handling").p("Information is labelled or recorded with its class, and handled according to the table above. Customer data is Restricted unless the customer agrees otherwise.")
d.h("Retention").ul(
    "Customer data is deleted within {data_after_contract} of a contract ending, or sooner when the customer asks.",
    "Security logs are kept for {log_retention}.",
    "Other records are kept for [[10.b1]] unless the law requires longer.")
d.h("Disposal").p("Data is deleted so it cannot be recovered. Devices and media are wiped or destroyed by an approved provider, and a certificate of destruction is kept.")
d.fields = {"10.b1": f("default retention", "How long are other business records kept?", "", "seven years")}
DOCS.append(d)

d = Doc("11", "Security_Awareness_and_Training_Policy", "Security Awareness and Training Policy", "CC1.4, CC2.2",
        "Background checks, security training at joining and every year, and the training record. Filled from Training.")
d.h("Hiring").p("Before someone joins, {company} checks their identity, right to work and references, and carries out a background check where the law allows, in proportion to the role.")
d.h("Training").ul(
    "Everyone completes security awareness training within [[11.b1]] of joining and {training_cycle} after that.",
    "Engineers also complete secure development training.",
    "Phishing simulations are run at least [[11.b2]], and people who click are offered extra training.")
d.h("Record").source("training", ["Person", "Course", "Audience", "Completed", "Next due", "Result"])
d.fields = {
    "11.b1": f("onboarding training deadline", "How soon must new starters finish training?", "", "30 days"),
    "11.b2": f("phishing test frequency", "How often are phishing simulations run?", "", "twice a year"),
}
DOCS.append(d)

d = Doc("12", "Monitoring_and_Internal_Review_Procedure", "Monitoring and Internal Review Procedure", "CC4.1, CC4.2, CC5.1",
        "How controls are checked between audits, and how weaknesses are reported and fixed. Filled from Internal audits and Exceptions.")
d.h("Purpose").p("{company} checks that its controls keep working between SOC examinations, so that weaknesses are found and fixed by us and not reported by the auditor.")
d.h("Ongoing monitoring").ul(
    "Each control owner checks their own controls each time they run, and keeps the evidence.",
    "The {security_owner} tests a sample of controls every [[12.b1]] and records each result in Offset Attest.",
    "Automated monitoring alerts on configuration drift, failed backups and missing endpoint protection.")
d.h("Internal reviews").p("At least once a year, someone independent of the control owners reviews the programme: [[12.b2]]. The review is planned in advance and its results reported to the {approver}.").source(
    "audits", ["Review", "Scope", "Led by", "Date", "Status"])
d.h("Reporting and fixing weaknesses").ul(
    "Every failed test, review finding or auditor exception is recorded with an owner and a due date.",
    "Open items are reported to the {approver} at least every quarter.",
    "An item is closed only when the fix has been made and the control tested again.").source(
    "findings", ["Ref", "Raised", "Finding", "Owner", "Status"])
d.fields = {
    "12.b1": f("self-test frequency", "How often are controls sample-tested internally?", "", "quarter"),
    "12.b2": f("who reviews", "Who carries out the independent review?", "For example an internal auditor, or an outside consultant.", "an outside consultant"),
}
DOCS.append(d)

d = Doc("00", "System_Description", "System Description", "SOC 2 description criteria",
        "The description of the system your CPA firm reviews and publishes in the report. Filled from Scope, Assets and Vendors.")
d.h("Company and services").p("{company} provides {system_name}. [[00.b1]]")
d.h("The system and its boundaries").p("[[sys.boundaries]]")
d.h("Service commitments and system requirements").p("[[sys.commitments]]")
d.h("Components").p("Infrastructure and software inside the boundary, from the asset register:").source("assets", ["System", "Type", "Owner", "Classification", "Criticality", "Location"])
d.h("People").p("The {approver} has overall responsibility. The {security_owner} runs the security programme. [[00.b2]]")
d.h("Subservice organisations").p("{system_name} relies on the vendors below. Unless stated otherwise their controls are carved out of this report, and {company} monitors them by reviewing their own SOC reports.").source("vendors", ["Service", "Vendor", "Assurance"])
d.h("Complementary user entity controls").p("Customers are responsible for:").ul(
    "managing which of their own people have access to {system_name}, and removing access when they leave;",
    "keeping their own credentials secret and using multi-factor authentication where offered;",
    "reviewing the data they send to the service for accuracy;",
    "telling {company} promptly of any suspected security incident at {security_contact}.")
d.fields = {
    "00.b1": f("service summary", "What does the service do, and for whom?", "Two or three sentences a customer would recognise.", "", True),
    "sys.boundaries": f("system and its boundaries", "The system and its boundaries", "Taken from Scope. Type here to use different wording.", "", True),
    "sys.commitments": f("service commitments and system requirements", "Service commitments and system requirements", "Taken from Scope. Type here to use different wording.", "", True),
    "00.b2": f("teams", "Which teams run the system?", "For example engineering, operations, customer support.", "", True),
}
DOCS.append(d)


# ── writing Word files ───────────────────────────────────────────────────────

def blank_run(par, tag, placeholder):
    """A yellow <placeholder> in a content control tagged with the blank's name."""
    run = par.add_run(f"<{placeholder}>")
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    r = run._r
    sdt, pr, content = OxmlElement("w:sdt"), OxmlElement("w:sdtPr"), OxmlElement("w:sdtContent")
    alias, t = OxmlElement("w:alias"), OxmlElement("w:tag")
    alias.set(qn("w:val"), placeholder[:60])
    t.set(qn("w:val"), tag)
    pr.append(alias)
    pr.append(t)
    sdt.append(pr)
    sdt.append(content)
    r.addprevious(sdt)
    content.append(r)


def write(par, text, doc, used):
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        shared, own = m.group(1), m.group(2)
        if shared:
            if shared.startswith("$"):
                blank_run(par, shared, SPECIAL[shared])
            else:
                if shared not in SHARED:
                    raise SystemExit(f"{doc.file}: unknown shared answer {shared}")
                used.add(shared)
                blank_run(par, shared, SHARED[shared][1])
        else:
            if own not in doc.fields:
                raise SystemExit(f"{doc.file}: no field defined for {own}")
            blank_run(par, own, doc.fields[own]["ph"])
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def plain(text):
    """The sentence a blank sits in, with the blanks shown as their placeholders."""
    def sub(m):
        if m.group(1):
            k = m.group(1)
            return f"<{SPECIAL[k]}>" if k.startswith("$") else f"<{SHARED[k][1]}>"
        return "..."
    return TOKEN.sub(sub, text)


def build(doc):
    w = Document()
    st = w.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    used = set()

    # Instructions: everything before the first page break, dropped when filled in.
    w.add_heading("How to use this template", level=2)
    w.add_paragraph(
        "Yellow <placeholders> are where your own details go. Fill them in on the Templates tab in "
        "Offset Attest and download the document filled in, or edit it here in Word.")
    w.add_paragraph(
        "Change anything that does not match what you really do. Your CPA firm will test what this "
        "document says, so a promise you do not keep becomes an exception in your SOC report.")
    w.add_paragraph(f"Relevant criteria: {doc.ref}.")
    w.add_page_break()

    # Header: whose document it is.
    hp = w.sections[0].header.paragraphs[0]
    write(hp, "{company}", doc, used)
    hp.add_run(f"  |  {doc.title}")

    title = w.add_heading(level=0)
    write(title, "{company} " + doc.title, doc, used)
    info = w.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    for label, key in [("Owner", "$owner"), ("Approver", "$approver"), ("Version", "$version"),
                       ("Effective date", "$approval_date"), ("Next review", "$next_review"),
                       ("Classification", "$classification")]:
        row = info.add_row().cells
        row[0].text = label
        write(row[1].paragraphs[0], "{" + key + "}", doc, used)
    w.add_paragraph()

    items, contexts = [], {}
    for kind, payload, section in doc.body:
        if kind == "h":
            w.add_heading(payload, level=1)
            continue
        if kind in ("p", "li"):
            par = w.add_paragraph(style="List Bullet" if kind == "li" else None)
            write(par, payload, doc, used)
            for m in TOKEN.finditer(payload):
                if m.group(2):
                    contexts.setdefault(m.group(2), plain(payload))
                    items.append({"t": "field", "section": section, "name": m.group(2)})
        elif kind == "table":
            header, rows = payload
            t = w.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            for i, h in enumerate(header):
                t.rows[0].cells[i].text = h
            for r in rows:
                cells = t.add_row().cells
                for i, c in enumerate(r):
                    write(cells[i].paragraphs[0], c, doc, used)
                    for m in TOKEN.finditer(c):
                        if m.group(2):
                            contexts.setdefault(m.group(2), plain(" | ".join(r)))
                            items.append({"t": "field", "section": section, "name": m.group(2)})
            w.add_paragraph()
        elif kind == "source":
            name, header = payload
            t = w.add_table(rows=2, cols=len(header))
            t.style = "Table Grid"
            cap = OxmlElement("w:tblCaption")
            cap.set(qn("w:val"), f"source:{name}")
            t._tbl.tblPr.append(cap)
            for i, h in enumerate(header):
                t.rows[0].cells[i].text = h
                run = t.rows[1].cells[i].paragraphs[0].add_run(f"<{h}>")
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            items.append({"t": "source", "section": section, "source": name, "header": header})
            w.add_paragraph()

    for k in doc.fields:
        if k not in contexts:
            raise SystemExit(f"{doc.file}: field {k} is defined but not used")
    w.save(OUT / doc.file)
    return {
        "file": doc.file, "no": doc.no, "title": doc.title, "id": doc.ref, "items": items,
        "fields": {k: {"label": v["label"], "hint": v["hint"], "default": v["default"],
                       "long": v["long"], "context": contexts[k]} for k, v in doc.fields.items()},
        "shared": sorted(used),
    }


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for old in OUT.glob("*.docx"):
        old.unlink()
    DOCS.sort(key=lambda d: d.no)
    documents = [build(d) for d in DOCS]
    manifest = {
        "version": 1, "ownerKey": "security_owner", "approverKey": "approver", "groups": GROUPS,
        "shared": [{"name": k, "group": g, "label": label, **({"hint": hint} if hint else {}),
                    **({"default": default} if default else {}), **({"long": True} if long else {})}
                   for k, (g, _ph, label, hint, default, long) in SHARED.items()],
        "documents": documents,
    }
    (OUT / "fields.json").write_text(json.dumps(manifest, indent=1, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n")
    index = {
        "note": "Starter policies for a SOC 2 report, written in plain language. Each one says only what most "
                "service organisations do; change it to match what you really do before it is approved.",
        "documents": [{"file": d.file, "title": d.title, "about": d.about} for d in DOCS],
    }
    (OUT / "index.json").write_text(json.dumps(index, indent=2, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n")
    print(f"{len(DOCS)} templates written to {OUT}")


if __name__ == "__main__":
    main()
